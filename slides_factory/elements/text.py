"""Text element — a titled/body/bulleted text box placed in a cell."""

from __future__ import annotations

from pptx.enum.text import MSO_ANCHOR
from pptx.slide import Slide
from pydantic import BaseModel

from slides_factory.elements.base import Box, style_paragraph
from slides_factory.layout.fonts import apply_shape_font
from slides_factory.render_context import RenderContext
from slides_factory.styling import theme
from slides_factory.styling.models import TextStyle


class TextProps(BaseModel):
    """Content props for the text element."""

    text: str = ""
    bullets: list[str] = []


def render_text(
    slide: Slide,
    box: Box,
    props: TextProps,
    style: TextStyle,
    ctx: RenderContext,
) -> None:
    """Render text and optional bullet lines."""
    left, top, width, height = box
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP

    body_size = theme.font_size_pt(style.text_size)

    lines: list[tuple[str, bool]] = []
    if props.text:
        lines.append((props.text, False))
    for bullet in props.bullets:
        lines.append((f"\u2022 {bullet}", True))
    if not lines:
        lines.append(("", False))

    for index, (content, _is_bullet) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = content
        style_paragraph(
            paragraph,
            ctx,
            size_pt=body_size,
            bold=style.bold,
            color_token=style.text_color,
            align=style.align,
        )

    apply_shape_font(textbox, ctx, style.font)
